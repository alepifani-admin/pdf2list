import csv
from pathlib import Path

from openpyxl import load_workbook

from pdf2list.core.export import export_configs_csv, export_csv, export_xlsx
from pdf2list.core.models import Column, ExtractionResult, Row, Table
from pdf2list.core.settings import AppSettings


def sample_result() -> ExtractionResult:
    cols = [
        Column("Descrizione", 0, 100, kind="text"),
        Column("Codice", 0, 0, kind="text"),
        Column("Kg", 0, 0, kind="number"),
        Column("Prezzo", 0, 0, kind="price"),
    ]
    rows = [
        Row(["FIANCHI", "", "", ""], 1, kind="group", group="FIANCHI"),
        Row(["Fianco", "COB90", "3,1", "€ 97,00"], 1, group="FIANCHI"),
        Row(["Ripiano", "COBR90", "2,7", "56,50"], 2, group="RIPIANI"),
    ]
    table = Table(columns=cols, rows=rows, pages=[1, 2],
                  title="COBA - TABELLA PREZZI E COMPONENTI")
    cards = Table(
        columns=[Column("Prodotto", 0, 0), Column("Codice", 0, 0),
                 Column("Prezzo", 0, 0, kind="price")],
        rows=[Row(["JOLLY", "JOL", "52,00"], 3)],
        pages=[3], title="Schede prodotto", source="cards",
    )
    return ExtractionResult("source.pdf", [table, cards], 37,
                            meta={"configs": [
                                {"sezione": "SISTEMI", "prodotto": "COBA",
                                 "codice": "2xCOB90", "prezzo": "819,00",
                                 "pagina": 6, "note": "Codice Sistema"}
                            ]})


DB_HEADERS = ["Tipo", "Tabella", "Gruppo", "Prodotto", "Descrizione",
              "Codice", "Prezzo", "Kg", "Note", "Pagina"]


def test_export_xlsx_uses_one_database_sheet(tmp_path: Path):
    result = sample_result()
    path = tmp_path / "out.xlsx"
    export_xlsx(result, path, AppSettings())
    wb = load_workbook(path)

    assert wb.sheetnames == ["Listino"]
    ws = wb["Listino"]
    assert [c.value for c in ws[1]] == DB_HEADERS
    assert ws.max_row == 5  # header + 2 articles + 1 card + 1 config
    assert "ListinoDB" in ws.tables
    assert ws.tables["ListinoDB"].autoFilter is not None
    assert ws.freeze_panes == "A2"

    rows = [r for r in range(2, ws.max_row + 1)]
    fianco = next(r for r in rows
                  if ws.cell(row=r, column=5).value == "Fianco")
    assert ws.cell(row=fianco, column=1).value == "Articolo"
    assert ws.cell(row=fianco, column=2).value == "COBA"
    assert ws.cell(row=fianco, column=3).value == "FIANCHI"
    assert ws.cell(row=fianco, column=5).value == "Fianco"
    assert ws.cell(row=fianco, column=6).value == "COB90"
    assert ws.cell(row=fianco, column=7).value == 97.0
    assert ws.cell(row=fianco, column=8).value == 3.1
    assert ws.cell(row=fianco, column=10).value == 1

    configs = [r for r in range(2, ws.max_row + 1)
               if ws.cell(row=r, column=1).value == "Configurazione"]
    assert len(configs) == 1
    cfg = configs[0]
    assert ws.cell(row=cfg, column=2).value == "SISTEMI"
    assert ws.cell(row=cfg, column=4).value == "COBA"
    assert ws.cell(row=cfg, column=6).value == "2xCOB90"
    assert ws.cell(row=cfg, column=7).value == 819.0
    assert ws.cell(row=cfg, column=9).value == "Codice Sistema"
    assert ws.cell(row=cfg, column=10).value == 6


def test_database_has_group_column_instead_of_group_rows(tmp_path: Path):
    result = sample_result()
    path = tmp_path / "groups.xlsx"
    export_xlsx(result, path, AppSettings())
    ws = load_workbook(path)["Listino"]
    groups = [ws.cell(row=r, column=3).value for r in range(2, ws.max_row + 1)]
    assert "FIANCHI" in groups
    assert "RIPIANI" in groups
    # group headings are data, not merged full-width rows
    assert not ws.merged_cells.ranges


def test_database_unions_heterogeneous_tables_and_keeps_numbers(tmp_path: Path):
    first = Table(
        columns=[Column("Codice", 0, 0), Column("Descrizione", 0, 0),
                 Column("Prezzo", 0, 0, kind="price"),
                 Column("L", 0, 0, kind="number"),
                 Column("H", 0, 0, kind="number")],
        rows=[Row(["A1", "Pannello", "10,00", "60", "80"], 1)],
        pages=[1], title="Prima tabella",
    )
    second = Table(
        columns=[Column("Codice", 0, 0), Column("Descrizione", 0, 0),
                 Column("Prezzo", 0, 0, kind="price"),
                 Column("Prezzo 2", 0, 0, kind="price")],
        rows=[Row(["B2", "Base", "20,00", "25,00"], 2)],
        pages=[2], title="Seconda tabella",
    )
    result = ExtractionResult("source.pdf", [first, second], 2)
    path = tmp_path / "union.xlsx"
    export_xlsx(result, path, AppSettings())
    wb = load_workbook(path)
    assert wb.sheetnames == ["Listino"]
    ws = wb["Listino"]
    headers = [c.value for c in ws[1]]
    assert headers.index("Prezzo") + 1 == headers.index("Prezzo 2")
    assert "L" in headers and "H" in headers
    values = {ws.cell(row=r, column=1).value: r for r in range(2, ws.max_row + 1)}
    assert ws.cell(row=2, column=headers.index("L") + 1).value == 60.0
    assert ws.cell(row=2, column=headers.index("H") + 1).value == 80.0
    assert ws.cell(row=3, column=headers.index("Prezzo 2") + 1).value == 25.0


def test_export_xlsx_empty_result_still_has_one_sheet(tmp_path: Path):
    result = ExtractionResult("empty.pdf", [], 1)
    path = tmp_path / "empty.xlsx"
    export_xlsx(result, path, AppSettings())
    wb = load_workbook(path)
    assert wb.sheetnames == ["Listino"]
    ws = wb["Listino"]
    assert ws.max_row == 1
    assert [c.value for c in ws[1]][:7] == [
        "Tipo", "Tabella", "Gruppo", "Prodotto", "Descrizione", "Codice",
        "Prezzo",
    ]
    assert "ListinoDB" not in ws.tables


def test_export_csv_matches_database_columns(tmp_path: Path):
    result = sample_result()
    path = tmp_path / "out.csv"
    export_csv(result, path, AppSettings())
    text = path.read_text(encoding="utf-8-sig")
    rows = list(csv.reader(text.splitlines(), delimiter=";"))
    assert rows[0] == DB_HEADERS
    assert len(rows) == 5
    assert rows[1] == ["Articolo", "COBA", "FIANCHI", "", "Fianco", "COB90",
                       "97,00", "3,1", "", "1"]
    assert rows[2][0:3] == ["Articolo", "COBA", "RIPIANI"]
    assert rows[3][0:6] == ["Articolo", "Schede prodotto", "", "JOLLY",
                            "", "JOL"]
    assert rows[4] == ["Configurazione", "SISTEMI", "", "COBA", "",
                       "2xCOB90", "819,00", "", "Codice Sistema", "6"]


def test_export_configs_csv(tmp_path: Path):
    result = sample_result()
    path = tmp_path / "config.csv"
    out = export_configs_csv(result, path, AppSettings())
    assert out == path
    text = path.read_text(encoding="utf-8-sig")
    assert "Codice configurazione" in text
    assert "2xCOB90" in text
    assert "819,00" in text


def test_table_title_cannot_create_an_extra_sheet(tmp_path: Path):
    table = Table(
        columns=[Column("Codice", 0, 0), Column("Prezzo", 0, 0, kind="price")],
        rows=[Row(["A1", "10,00"], 1)], pages=[1], title="Configurazioni",
    )
    result = ExtractionResult("source.pdf", [table], 1)
    path = tmp_path / "c.xlsx"
    export_xlsx(result, path, AppSettings())
    wb = load_workbook(path)
    assert wb.sheetnames == ["Listino"]


def test_formula_injection_is_neutralised(tmp_path: Path):
    table = Table(
        columns=[Column("Descrizione", 0, 0), Column("Codice", 0, 0),
                 Column("Prezzo", 0, 0, kind="price")],
        rows=[Row(["=1+1", "-2+3", "10,00"], 1)], pages=[1], title="=1+1",
    )
    result = ExtractionResult("source.pdf", [table], 1, meta={"configs": [
        {"sezione": "@x", "prodotto": "P", "codice": "=1+1",
         "prezzo": "10,00", "pagina": 1, "note": "=cmd"}
    ]})
    xlsx = tmp_path / "i.xlsx"
    export_xlsx(result, xlsx, AppSettings())
    wb = load_workbook(xlsx)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                assert not (isinstance(cell.value, str) and cell.value.startswith("="))

    csv_path = tmp_path / "i.csv"
    export_csv(result, csv_path, AppSettings())
    rows = list(csv.reader(csv_path.read_text(encoding="utf-8-sig").splitlines(),
                           delimiter=";"))
    headers = rows[0]
    desc = headers.index("Descrizione")
    code = headers.index("Codice")
    assert rows[1][desc].startswith("'")
    assert rows[1][code].startswith("'")
    assert rows[2][code].startswith("'")

    cfg_path = tmp_path / "i-config.csv"
    export_configs_csv(result, cfg_path, AppSettings())
    cfg = list(csv.reader(cfg_path.read_text(encoding="utf-8-sig").splitlines(),
                          delimiter=";"))
    assert cfg[1][0].startswith("'")
    assert cfg[1][2].startswith("'")
