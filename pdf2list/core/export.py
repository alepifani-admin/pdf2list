"""Excel and CSV export as a single, database-ready table."""

from __future__ import annotations

import csv
import logging
import re
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.filters import AutoFilter
from openpyxl.worksheet.table import Table as ExcelTable
from openpyxl.worksheet.table import TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet

from .. import __app_name__, __version__
from .layout import _short_title
from .models import ExtractionResult, Table
from .normalize import parse_number, safe_filename
from .settings import AppSettings

log = logging.getLogger(__name__)

DARK = "2E3033"
GREY = "F2F3F5"
WHITE = "FFFFFF"
BORDER = "BFC4C8"

# Identity columns always present in the database sheet, in this order.
CORE_COLUMNS = ["Tipo", "Tabella", "Gruppo", "Prodotto", "Descrizione", "Codice"]
RESERVED_COLUMNS = set(CORE_COLUMNS) | {"Prezzo", "Note", "Pagina"}
# Optional columns are ordered with the most useful ones first.
PREFERRED = ["Struttura", "Componenti", "Pannelli", "Ganci", "Coperchio",
             "Ripiani", "Formato", "Elementi", "Misure", "Imballaggio",
             "Kg", "Quantità", "Qty", "Sezione", "Sottosezione"]

DB_SHEET = "Listino"
DB_TABLE_NAME = "ListinoDB"
DB_TABLE_STYLE = "TableStyleMedium2"


@dataclass
class ExportResult:
    xlsx: Path | None = None
    csv: Path | None = None


def _fill(color: str) -> PatternFill:
    return PatternFill("solid", fgColor=color)


def _style_header(ws: Worksheet, row: int, ncols: int, height: int = 30) -> None:
    thin = Side(style="thin", color=BORDER)
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = _fill(DARK)
        cell.font = Font(bold=True, color=WHITE, size=10)
        cell.alignment = Alignment(horizontal="center", vertical="center",
                                   wrap_text=True)
        cell.border = Border(bottom=thin)
    ws.row_dimensions[row].height = height


def _set_widths(ws: Worksheet, widths: list[int | float]) -> None:
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _auto_widths(ws: Worksheet, ncols: int, rows: int,
                 headers: list[str]) -> None:
    widths: list[float] = []
    for ci in range(1, ncols + 1):
        header = headers[ci - 1] if ci - 1 < len(headers) else ""
        longest = max(8, min(38, len(str(header)) + 2))
        for r in range(2, min(rows, 220) + 1):
            value = ws.cell(row=r, column=ci).value
            if value is None:
                continue
            text = str(value)
            longest = max(longest, min(48, len(text) + 2))
        widths.append(longest)
    _set_widths(ws, widths)


def _is_price(col_name: str, kind: str = "") -> bool:
    if kind == "price":
        return True
    key = (col_name or "").strip().lower()
    if re.match(r"^cat\.?\s*[0-9a-z]$", key) or key.startswith("stamp"):
        return True
    return any(x in key for x in ("prezzo", "price", "prix", "precio", "preis",
                                  "importo", "€", "eur"))


def _price_key(name: str) -> tuple[int, int | str]:
    match = re.match(r"(?i)^prezzo(?:\s*(\d+))?$", (name or "").strip())
    if match:
        return (0, int(match.group(1) or 1))
    return (1, (name or "").lower())


def _money_format(symbol: str) -> str:
    sym = "".join(c for c in (symbol or "") if c not in '"&').strip()[:3]
    return f'#,##0.00 "{sym}"' if sym else "#,##0.00"


def _protect_formula(text: str) -> str:
    """Neutralise spreadsheet formula injection in extracted text."""
    if not text:
        return text
    if text[0] in "\t\r":
        return "'" + text
    if text[0] in "=+-@" and len(text) > 1:
        return "'" + text
    return text


def _prepare_cell(value: str, kind: str, currency: str):
    """Return (value, number_format) for a cell."""
    if value is None or str(value).strip() == "":
        return None, None
    text = str(value).strip()
    if kind in {"price", "number"}:
        number = parse_number(text)
        if number is not None:
            if kind == "price":
                return number, _money_format(currency)
            return number, "0.0##"
    # avoid spreadsheet formula injection from extracted text
    return _protect_formula(text), None


# ---------------------------------------------------------------------------
# Database columns and rows
# ---------------------------------------------------------------------------

def _column_kinds(tables: list[Table]) -> dict[str, str]:
    """Merge column kinds across tables (price > number > text)."""
    rank = {"text": 0, "number": 1, "price": 2}
    kinds: dict[str, str] = {}
    for table in tables:
        for column in table.columns:
            name = (column.name or "").strip()
            if not name or name in CORE_COLUMNS or name in {"Note", "Pagina"}:
                continue
            kind = column.kind or "text"
            if _is_price(name, kind):
                kind = "price"
            if rank.get(kind, 0) >= rank.get(kinds.get(name, "text"), 0):
                kinds[name] = kind
    return kinds


def database_headers(tables: list[Table], include_page: bool = True
                     ) -> tuple[list[str], dict[str, str]]:
    """Stable database schema: identity, prices, extra fields, note, page."""
    kinds = _column_kinds(tables)
    prices = sorted((h for h in kinds if _is_price(h, kinds[h])), key=_price_key)
    if not prices:
        prices = ["Prezzo"]
    kinds.setdefault("Prezzo", "price")

    extra: list[str] = []
    for name in [h for h in PREFERRED if h in kinds] + list(kinds):
        if name not in extra and name not in RESERVED_COLUMNS:
            extra.append(name)

    headers = CORE_COLUMNS + prices + extra + ["Note"]
    if include_page:
        headers.append("Pagina")
    unique: list[str] = []
    for name in headers:
        if name not in unique:
            unique.append(name)
    return unique, kinds


def _table_record(table: Table, row, headers: list[str],
                  include_page: bool) -> dict:
    mapping: dict[str, str] = {}
    for index, name in enumerate(table.headers):
        if index < len(row.cells):
            mapping[name] = row.cells[index]
    record = {
        "Tipo": "Articolo",
        "Tabella": _short_title(table.title) or DB_SHEET,
        "Gruppo": row.group or mapping.get("Gruppo", ""),
        "Prodotto": mapping.get("Prodotto", ""),
        "Descrizione": mapping.get("Descrizione", ""),
        "Codice": mapping.get("Codice", ""),
        "Prezzo": mapping.get("Prezzo", ""),
        "Note": mapping.get("Note", ""),
    }
    for name in headers:
        if name not in record:
            record[name] = mapping.get(name, "")
    if include_page:
        record["Pagina"] = row.page
    return record


def _config_record(config: dict, headers: list[str],
                   include_page: bool) -> dict:
    record = {name: "" for name in headers}
    record["Tipo"] = "Configurazione"
    record["Tabella"] = config.get("sezione") or "Configurazioni"
    record["Prodotto"] = config.get("prodotto", "")
    record["Codice"] = config.get("codice", "")
    record["Prezzo"] = config.get("prezzo", "")
    record["Note"] = config.get("note", "")
    if include_page:
        record["Pagina"] = config.get("pagina", "")
    return record


def database_records(result: ExtractionResult, headers: list[str],
                     include_page: bool) -> list[dict]:
    records: list[dict] = []
    for table in result.tables:
        for row in table.data_rows():
            records.append(_table_record(table, row, headers, include_page))
    for config in result.meta.get("configs") or []:
        records.append(_config_record(config, headers, include_page))
    return records


def _write_record(ws: Worksheet, index: int, record: dict,
                  headers: list[str], kinds: dict[str, str],
                  currency: str) -> None:
    for column, name in enumerate(headers, start=1):
        raw = record.get(name, "")
        if name == "Pagina":
            if raw is None or str(raw).strip() == "":
                value = None
            else:
                try:
                    value = int(raw)
                except (TypeError, ValueError):
                    value = _protect_formula(str(raw))
            fmt = None
        else:
            value, fmt = _prepare_cell(raw, kinds.get(name, "text"), currency)
        cell = ws.cell(row=index, column=column, value=value)
        if fmt:
            cell.number_format = fmt
        right = bool(fmt and ("0.00" in fmt or "0.0" in fmt))
        cell.alignment = Alignment(
            vertical="top",
            wrap_text=name in {"Descrizione", "Note", "Tabella"} and bool(
                value and len(str(value)) > 32),
            horizontal="right" if right else "left",
        )


def _add_database_table(ws: Worksheet, headers: list[str], has_rows: bool) -> None:
    if not has_rows:
        return
    reference = f"A1:{get_column_letter(len(headers))}{ws.max_row}"
    try:
        table = ExcelTable(displayName=DB_TABLE_NAME, ref=reference)
        table.autoFilter = AutoFilter(ref=reference)
        table.tableStyleInfo = TableStyleInfo(
            name=DB_TABLE_STYLE, showFirstColumn=False, showLastColumn=False,
            showRowStripes=True, showColumnStripes=False,
        )
        ws.add_table(table)
    except Exception as exc:  # pragma: no cover - defensive fallback
        log.debug("cannot create Excel table: %s", exc)
        ws.auto_filter.ref = reference


# ---------------------------------------------------------------------------
# Workbook
# ---------------------------------------------------------------------------

def build_workbook(result: ExtractionResult, settings: AppSettings) -> Workbook:
    export = settings.export
    wb = Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet(DB_SHEET)

    tables = [t for t in result.tables if t.data_rows()]
    headers, kinds = database_headers(tables, export.include_page_column)
    records = database_records(result, headers, export.include_page_column)

    ws.append(headers)
    _style_header(ws, 1, len(headers))
    for index, record in enumerate(records, start=2):
        _write_record(ws, index, record, headers, kinds,
                      export.currency_symbol)

    if export.freeze_header:
        ws.freeze_panes = "A2"
    if export.auto_filter:
        _add_database_table(ws, headers, bool(records))
    _auto_widths(ws, len(headers), ws.max_row, headers)
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 90
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.oddFooter.left.text = f"{__app_name__} - {Path(result.source).stem}"
    ws.oddFooter.left.size = 8
    ws.oddFooter.right.text = "Pag. &P di &N"
    ws.oddFooter.right.size = 8
    ws.print_title_rows = "1:1" if records else None
    ws.sheet_properties.tabColor = DARK
    if not records:
        ws["A1"].comment = Comment(
            "Nessuna tabella riconosciuta nel PDF.", __app_name__
        )

    # Metadata lives in the file properties: the workbook keeps one data sheet.
    detected = result.meta.get("detected_format") or {}
    ocr_pages = result.meta.get("ocr_pages") or []
    props = wb.properties
    props.creator = f"{__app_name__} {__version__}"
    props.lastModifiedBy = __app_name__
    props.title = f"{Path(result.source).stem} - {__app_name__}"
    props.subject = str(detected.get("summary") or "Listino prezzi")
    props.keywords = "listino; prezzi; database; PDF2LIST"
    details = [
        f"File: {Path(result.source).name}",
        f"Pagine PDF: {result.pages}",
        f"Tabelle: {len(tables)}",
        f"Righe articoli: {sum(len(t.data_rows()) for t in tables)}",
        f"Configurazioni: {len(result.meta.get('configs') or [])}",
        f"Profilo: {result.meta.get('profile_id') or settings.profile_id}",
        f"Formato: {detected.get('summary') or '—'}",
        f"OCR: {'usato (' + str(len(ocr_pages)) + ' pagine)' if ocr_pages else 'non usato'}",
        f"AI: {'usata' if result.meta.get('ai_used') else 'non usata'}",
        f"Data: {datetime.now().strftime('%d/%m/%Y %H:%M')}",
        "I dati sono estratti automaticamente: verificare prezzi, codici e unità.",
    ]
    warnings = list(result.warnings) + [w for t in tables for w in t.warnings]
    if warnings:
        details.append("Avvisi: " + " | ".join(w for w in warnings if w)[:4000])
    props.description = "\n".join(details)

    wb.active = 0
    return wb


def export_xlsx(result: ExtractionResult, path: str | Path,
                settings: AppSettings) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = build_workbook(result, settings)
    wb.save(path)
    return path


# --------------------------------------------------------------------------
# CSV
# --------------------------------------------------------------------------

def _csv_value(value: str, kind: str, export) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if not text:
        return ""
    if kind in {"price", "number"}:
        num = parse_number(text)
        if num is not None:
            if kind == "price":
                out = f"{num:.2f}"
            else:
                out = f"{num:.3f}".rstrip("0").rstrip(".")
            return out.replace(".", ",") if export.decimal_comma else out
    return _protect_formula(text)


def export_csv(result: ExtractionResult, path: str | Path,
               settings: AppSettings) -> Path:
    export = settings.export
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    tables = [t for t in result.tables if t.data_rows()]
    headers, kinds = database_headers(tables, export.include_page_column)
    records = database_records(result, headers, export.include_page_column)
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f, delimiter=export.csv_delimiter or ";",
                            quoting=csv.QUOTE_MINIMAL)
        writer.writerow(headers)
        for record in records:
            cells: list[str] = []
            for name in headers:
                raw = record.get(name, "")
                if name == "Pagina":
                    cells.append("" if raw in (None, "") else str(raw))
                else:
                    cells.append(_csv_value(raw, kinds.get(name, "text"), export))
            writer.writerow(cells)
    return path


def export_configs_csv(result: ExtractionResult, path: str | Path,
                       settings: AppSettings) -> Path | None:
    configs = result.meta.get("configs") or []
    if not configs:
        return None
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    headers = ["Sezione", "Prodotto", "Codice configurazione", "Prezzo EUR",
               "Pagina PDF", "Dettaglio"]
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f, delimiter=settings.export.csv_delimiter or ";",
                            quoting=csv.QUOTE_MINIMAL)
        writer.writerow(headers)
        for c in configs:
            price = c.get("prezzo", "")
            if price:
                num = parse_number(price)
                if num is not None:
                    price = f"{num:.2f}".replace(
                        ".", "," if settings.export.decimal_comma else ".")
            writer.writerow([_protect_formula(c.get("sezione", "")),
                             _protect_formula(c.get("prodotto", "")),
                             _protect_formula(c.get("codice", "")), price,
                             c.get("pagina", ""), _protect_formula(c.get("note", ""))])
    return path


def default_output_paths(source: str | Path, out_dir: str | Path | None = None) -> tuple[Path, Path]:
    src = Path(source)
    base = safe_filename(src.stem, "listino")
    folder = Path(out_dir) if out_dir else src.parent
    return folder / f"{base}.xlsx", folder / f"{base}.csv"


def configs_path_for(csv_path: str | Path) -> Path:
    p = Path(csv_path)
    return p.with_name(f"{p.stem} - CONFIGURAZIONI.csv")
